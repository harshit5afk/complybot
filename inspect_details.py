

import os
import pptx

prs = pptx.Presentation(os.path.expanduser(r'~\Downloads\6aa7f1b39e957_infothon_7_0_template.pptx'))
print(f'Slide count: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides):
    print(f'Slide {idx+1}: {len(slide.shapes)} shapes')
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            first_p = s.text_frame.paragraphs[0]
            font_info = ''
            if first_p.runs:
                r = first_p.runs[0]
                font_info = f'font={r.font.name}, size={r.font.size.pt if r.font.size else "None"}, color={r.font.color.rgb if r.font.color and hasattr(r.font.color, "rgb") else "None"}'
            print(f'   [{s.name}] ({s.left},{s.top},{s.width},{s.height}): "{s.text_frame.text[:40].strip()}" {font_info}')
